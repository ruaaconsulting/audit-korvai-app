import hashlib
import csv
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pathlib import Path


def _flatten_outline(outline) -> list[str]:
    titles = []
    for item in outline:
        if isinstance(item, list):
            titles.extend(_flatten_outline(item))   # this is a nested branch - recurse into it
        elif hasattr(item, "title"):
            titles.append(item.title)
    return titles

def fingerprint(file_path: Path) -> str:
    return hashlib.sha256(file_path.read_bytes()).hexdigest()

def extract_skeleton_md(file_path: Path) -> dict:
    text = file_path.read_text(encoding="utf-8")
    headings = [line.strip() for line in text.split("\n") if line.startswith("#")]
    return {"document_id": file_path.stem, "structure": headings}

def extract_skeleton_pdf(file_path: Path) -> dict:
    reader = PdfReader(file_path)
    headings = _flatten_outline(reader.outline)
    if not headings:
        print(f"Warning: {file_path.name} has no bookmarks - skeleton will be empty")
    return {"document_id": file_path.stem, "structure": headings}

def extract_skeleton_docx(file_path: Path) -> dict:
    doc = Document(file_path)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    return {"document_id": file_path.stem, "structure": headings}

def extract_skeleton_xlsx(file_path: Path) -> dict:
    workbook = load_workbook(file_path, read_only=True)
    sheets = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        first_row = []
        for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            first_row = list(row)
            break
        sheets.append({"sheet_name": sheet_name, "columns": first_row})
    workbook.close()
    return {"document_id": file_path.stem, "sheets": sheets}

def extract_skeleton_csv(file_path: Path) -> dict:
    with open(file_path, mode="r", encoding="utf-8") as f:
        reader = csv.reader(f)
        header_row = next(reader)
    return {"document_id": file_path.stem, "sheets": [{"sheet_name": "default", "columns": header_row}]}


EXTRACTORS = {
    ".md": extract_skeleton_md,
    ".pdf": extract_skeleton_pdf,
    ".docx": extract_skeleton_docx,
    ".xlsx": extract_skeleton_xlsx,
    ".csv": extract_skeleton_csv,
}